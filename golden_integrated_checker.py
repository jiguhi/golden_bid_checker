# -*- coding: utf-8 -*-
"""황금이네 순위체커 + 승인 후 5% 입찰 조정. 예상입찰가 API 미사용."""
import base64
import hashlib
import hmac
import json
import os
import re
import time
import subprocess
import sys
from collections import Counter
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
from html.parser import HTMLParser
from pathlib import Path
from urllib.error import HTTPError, URLError
from urllib.parse import quote, unquote, urlencode, urlsplit
from urllib.request import Request, build_opener, HTTPRedirectHandler

BASE_DIR = Path(__file__).resolve().parent
PROFILE_DIR = BASE_DIR / "chrome_profile"
MAPPING_FILE = BASE_DIR / "material_ids.xlsx"
CONFIG_FILE = BASE_DIR / "naver_searchad_config.json"
AUTH_STATE_FILE = Path(os.environ.get("GOLDEN_AUTH_STATE_FILE", str(BASE_DIR / ".auth" / "naver_auth_state.json")))
RESULT_DIR = BASE_DIR / "results"
DEBUG_DIR = BASE_DIR / "debug"
JOURNAL_FILE = RESULT_DIR / "bid_history.jsonl"
TARGET_SELLER = "황금이네"
TARGET_STORE = "smartstore.naver.com/goldhouse"
KEYWORDS = [
    "흑염소진액", "도라지배즙", "호박즙", "CCA주스",
    "ABC주스", "다슬기", "석류즙", "흑마늘즙",
    "사과즙", "토마토즙", "비트즙", "다슬기즙", "강황가루",
    "여주즙", "개구리즙", "우슬닭발즙", "벌나무즙",
    "장어즙", "엉겅퀴즙", "칡즙", "붕어즙", "민들레즙", "헛개즙",
    "미나리즙", "야관문즙", "가물치즙", "잉어즙",
]
# 조회/결과/입찰 처리 순서: PC 전체 → 충분한 대기 → Mobile 전체.
# Search AD API 입찰 로직은 기존 그대로 유지합니다.
SURFACES = ("파워링크 PC", "쇼핑 PC", "파워링크 Mobile", "쇼핑 Mobile")
PAGE_SETTLE_MS = 1500
BETWEEN_KEYWORD_SECONDS = 3.0       # 같은 디바이스에서 검색 사이 고정 간격
BETWEEN_DEVICE_BATCH_SECONDS = 30.0  # PC 전체 완료 후 Mobile 시작 전 휴식
CDP_ENDPOINT = os.environ.get("GOLDEN_CDP_ENDPOINT", "http://127.0.0.1:9222")
BROWSER_MODE = os.environ.get("GOLDEN_BROWSER_MODE", "managed").strip().lower()
MAX_BID = 20000                 # 사용자 설정 상한, 매체 상한이 아님
MAX_RESULT_AGE_SECONDS = 1800   # 배치 조회 안정화로 실행이 길어져 30분 이상 지난 결과는 변경 금지
MIN_REBID_SECONDS = 600         # 같은 ID 재변경 간격: 10분
BID_REQUEST_ATTEMPTED = False


def clean(value):
    return " ".join(str(value or "").split())


def save_json(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding="utf-8")


def yes(prompt):
    return input(prompt + " [y 입력 시 진행 / Enter 취소] > ").strip().lower() == "y"


def acquire_lock():
    handle = (BASE_DIR / "golden_checker.lock").open("a+b")
    if handle.tell() == 0:
        handle.write(b"0")
        handle.flush()
    handle.seek(0)
    try:
        if os.name == "nt":
            import msvcrt
            msvcrt.locking(handle.fileno(), msvcrt.LK_NBLCK, 1)
        else:
            import fcntl
            fcntl.flock(handle.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
    except OSError:
        handle.close()
        raise RuntimeError("이 폴더의 프로그램이 이미 실행 중입니다.") from None
    return handle   # 종료 시 close하면 운영체제가 잠금을 해제합니다.


# page.content()를 분석하므로 이미지 지연 로딩과 화면 좌표에 의존하지 않습니다.
class Node:
    def __init__(self, tag="root", attrs=(), parent=None):
        self.tag, self.attrs, self.parent = tag, dict(attrs), parent
        self.parts = []

    def walk(self):
        for part in self.parts:
            if isinstance(part, Node):
                yield part
                yield from part.walk()

    def text(self):
        if self.tag in ("script", "style", "noscript"):
            return ""
        return "".join(p.text() if isinstance(p, Node) else p for p in self.parts)

    def cls(self, name):
        return name in (self.attrs.get("class") or "").split()

    def closest(self, tag):
        node = self
        while node is not None:
            if node.tag == tag:
                return node
            node = node.parent
        return None


class Document(HTMLParser):
    VOID = set("area base br col embed hr img input link meta param source track wbr".split())

    def __init__(self, source):
        super().__init__(convert_charrefs=True)
        self.root = Node()
        self.stack = [self.root]
        self.feed(source)
        self.close()

    def handle_starttag(self, tag, attrs):
        node = Node(tag, attrs, self.stack[-1])
        self.stack[-1].parts.append(node)
        if tag not in self.VOID:
            self.stack.append(node)

    def handle_endtag(self, tag):
        for index in range(len(self.stack) - 1, 0, -1):
            if self.stack[index].tag == tag:
                del self.stack[index:]
                return

    def handle_startendtag(self, tag, attrs):
        self.handle_starttag(tag, attrs)
        if tag not in self.VOID:
            self.handle_endtag(tag)

    def handle_data(self, data):
        self.stack[-1].parts.append(data)


def outcome(status, cards=None, reason=""):
    return {"status": status, "cards": cards or [], "reason": reason}


def ad_id(node):
    values = [node.attrs.get("data-slog-content", "")]
    for n in node.walk():
        values.extend(n.attrs.get(k, "") or "" for k in ("onclick", "href", "data-slog-content"))
    matches = set(re.findall(r"nad-[A-Za-z0-9-]+", unquote(" ".join(values))))
    return next(iter(matches)) if len(matches) == 1 else ""


def store_match(value):
    return bool(re.search(r"(?:https?://)?smartstore\.naver\.com/goldhouse(?:[/?#\s]|$)", unquote(value or ""), re.I))


def has_ad_plus_svg(children):
    # 실제 저장 HTML의 광고+ 버튼: 접근성 텍스트는 광고, +는 SVG path.
    signature = "M297.5v-3h-1v3h-3v1h3v3h1v-3h3v-1z"
    return any(
        re.sub(r"[\s,]+", "", p.attrs.get("d", "")) == signature
        for n in children if n.tag == "button" and clean(n.text()) == "광고"
        for p in n.walk() if p.tag == "path"
    )


def filter_shopping_cards(cards):
    """사용자 지정: 광고+ 카드는 순위에서 제외하고 나머지를 재번호화."""
    kept = [dict(c) for c in cards if not c.get("ad_plus")]
    for rank, card in enumerate(kept, 1):
        card["rank"] = rank
    result = outcome("확인", kept) if kept else outcome("판독불가", reason="광고+ 제외 후 일반 광고 없음")
    result["excluded_ad_plus"] = [c for c in cards if c.get("ad_plus")]
    return result


def parse_shopping(root, mobile):
    sections = [n for n in root.walk() if n.tag == "section" and n.attrs.get("data-slog-container") == "shp_tli"]
    if not sections and not mobile:
        titles = [n for n in root.walk() if n.tag in ("h2", "h3", "span", "strong", "a") and clean(n.text()) in ("네이버 가격비교", "네이버 쇼핑 인기상품")]
        sections = list(dict.fromkeys(n.closest("section") for n in titles if n.closest("section")))
    if len(sections) != 1:
        return outcome("판독불가", reason="쇼핑 영역 없음 또는 여러 영역")
    section = sections[0]
    cameras = [n for n in section.walk() if n.tag == "ul" and n.cls("flicking-camera")]
    scope = section
    if mobile and cameras:
        panels = [n for n in cameras[0].parts if isinstance(n, Node) and n.tag == "li"]
        if not panels:
            return outcome("판독불가", reason="첫 페이지 없음")
        scope = panels[0]     # 다음 페이지가 미리 로딩되어 있어도 포함하지 않음
    elif not mobile:
        # PC의 flicking-camera는 상품 목록이 아니라 브랜드 필터입니다.
        # 실제 상품 카드는 형제 ul 안의 긴 텍스트 li로 렌더링됩니다.
        lists = [n for n in section.walk() if n.tag == "ul"]
        card_list = next((n for n in lists if any(
            isinstance(child, Node)
            and child.tag == "li"
            and (
                len(clean(child.text())) >= 40
                or "광고" in clean(child.text())
                or any((desc.attrs.get("id") or "").startswith("view_type_guide_") for desc in child.walk())
            )
            for child in n.parts
        )), None)
        if card_list is not None:
            scope = card_list
    elif mobile:
        return outcome("판독불가", reason="모바일 페이지 구조 변경")
    guides = [n for n in scope.walk() if (n.attrs.get("id") or "").startswith("view_type_guide_")]
    nodes = [n.closest("li") for n in guides]
    if not nodes and not mobile:
        nodes = [n.closest("li") for n in scope.walk() if re.fullmatch(r"광고(?:\s*[+＋])?", clean(n.text()))]
    nodes = list(dict.fromkeys(n for n in nodes if n is not None))
    cards, seen = [], set()
    for node in nodes:
        children = list(node.walk())
        if not any(clean(n.text()).startswith("광고") for n in children):
            continue
        aid = ad_id(node)
        guide = next((n for n in children if (n.attrs.get("id") or "").startswith("view_type_guide_")), None)
        pid = guide.attrs["id"].replace("view_type_guide_", "") if guide else ""
        seller_nodes = [n for n in children if n.cls("PtxugWXH")]
        seller = clean(seller_nodes[0].text()) if seller_nodes else ""
        if not seller and not mobile:
            seller = next((TARGET_SELLER for n in children if clean(n.text()) == TARGET_SELLER or store_match(n.attrs.get("href")) or store_match(n.attrs.get("onclick")) or store_match(n.text())), "")
        if not seller and not mobile:
            for child in children:
                text = clean(child.text())
                match = re.match(r"^(.{1,30}?)광고(?:\s*공식)?", text)
                if match and clean(match.group(1)):
                    seller = clean(match.group(1))
                    break
        title = next((clean(n.text()) for n in children if n.tag == "strong" and len(clean(n.text())) > 3), "")
        identity = aid or pid
        if not identity and (mobile or not seller and not title):
            return outcome("판독불가", reason="상품/소재 ID 확인 실패")
        if identity and identity in seen:
            continue
        if identity:
            seen.add(identity)
        plus = has_ad_plus_svg(children) or any(re.fullmatch(r"광고\s*[+＋]", clean(value)) for n in children
                   for value in (n.text(), n.attrs.get("aria-label", ""), n.attrs.get("title", "")))
        cards.append({"rank": len(cards) + 1, "ad_id": aid, "product_id": pid, "seller": seller, "title": title, "ad_plus": plus})
    return filter_shopping_cards(cards) if cards else outcome("판독불가", reason="광고 카드 없음/로딩 미완료")


def parse_shopping_pc_live(page):
    """PC 쇼핑은 현재 DOM에서 카드 순서를 직접 읽습니다.

    PC 네이버 쇼핑은 같은 화면에서도 판매자 텍스트가 일반 요소로
    렌더링되거나 쇼핑 브리지 링크로만 남는 경우가 있어, HTMLParser
    결과만으로는 정상 카드가 판독불가가 될 수 있습니다.
    """
    try:
        data = page.evaluate(r"""
() => {
    const norm = value => String(value || "").replace(/\s+/g, " ").trim();
    const sectionTitle = value => {
        const text = norm(value);
        return text === "네이버 가격비교" || text === "네이버 쇼핑 인기상품";
    };

    const headings = Array.from(
        document.querySelectorAll("h1,h2,h3,h4,span,strong,a,div")
    );
    const heading = headings.find(node => sectionTitle(node.innerText || node.textContent));
    let section = heading ? heading.closest("section") : null;

    if (!section) {
        section = Array.from(document.querySelectorAll("section")).find(node => {
            const text = norm(node.innerText || node.textContent);
            return text.includes("네이버 가격비교") || text.includes("네이버 쇼핑 인기상품");
        }) || null;
    }

    if (!section) {
        return {ok: false, reason: "PC 쇼핑 영역 없음"};
    }

    // PC의 flicking-camera는 브랜드 필터이고, 상품 카드는 별도 ul입니다.
    const lists = Array.from(section.querySelectorAll("ul"));
    const cardList = lists.find(list => {
        const items = Array.from(list.children)
            .filter(node => node.tagName === "LI");
        return items.some(item => {
            const text = norm(item.innerText || item.textContent);
            return (
                text.length >= 40
                || text.includes("광고")
                || item.querySelector('[id^="view_type_guide_"]')
            );
        });
    });
    const scope = cardList || section;

    const labels = Array.from(scope.querySelectorAll("*")).filter(node => {
        const text = norm(node.innerText || node.textContent);
        return text.startsWith("광고");
    });

    const cards = [];
    const seen = new Set();

    for (const label of labels) {
        const card = label.closest("li");
        if (!card || seen.has(card)) continue;
        seen.add(card);

        const descendants = [card, ...card.querySelectorAll("*")];
        const values = [];

        for (const node of descendants) {
            for (const name of ["id", "href", "onclick", "data-slog-content"]) {
                const value = node.getAttribute && node.getAttribute(name);
                if (value) values.push(value);
            }
        }

        const ids = Array.from(new Set(values.join(" ").match(/nad-[A-Za-z0-9-]+/g) || []));
        const productIds = Array.from(new Set(
            descendants
                .map(node => node.id || "")
                .filter(value => value.startsWith("view_type_guide_"))
                .map(value => value.replace("view_type_guide_", ""))
        ));

        const cardText = norm(card.innerText || card.textContent);
        const targetText = cardText.includes("황금이네");
        const targetStore = Array.from(card.querySelectorAll("a")).some(anchor => {
            const values = [anchor.href || "", anchor.getAttribute("href") || "", anchor.getAttribute("onclick") || ""];
            return values.some(value => /smartstore\.naver\.com\/goldhouse/i.test(value));
        });

        const targetSeller = targetText || targetStore;
        const title = (
            Array.from(card.querySelectorAll("img[alt],strong"))
                .map(node => norm(node.alt || node.innerText || node.textContent))
                .find(value => value.length > 3) || ""
        );

        cards.push({
            ad_plus: descendants.some(node => [node.textContent, node.getAttribute("aria-label"), node.getAttribute("title")].some(value => /^광고\s*[+＋]$/.test(norm(value))))
                || descendants.some(node => node.tagName === "BUTTON" && norm(node.textContent) === "광고"
                    && Array.from(node.querySelectorAll("svg path")).some(path => (path.getAttribute("d") || "").replace(/[\s,]+/g, "") === "M297.5v-3h-1v3h-3v1h3v3h1v-3h3v-1z")),
            ad_id: ids.length === 1 ? ids[0] : "",
            product_id: productIds.length === 1 ? productIds[0] : "",
            seller: targetSeller ? "황금이네" : "",
            title,
            top: card.getBoundingClientRect().top,
            left: card.getBoundingClientRect().left,
        });
    }

    // PC 쇼핑은 DOM 순서와 화면 순서가 달라지는 경우가 있어
    // 실제 화면의 위→아래, 왼쪽→오른쪽 순서로 정렬합니다.
    cards.sort((a, b) => {
        const sameRow = Math.abs(a.top - b.top) < 20;
        return sameRow ? a.left - b.left : a.top - b.top;
    });

    return {ok: cards.length > 0, reason: cards.length ? "" : "PC 쇼핑 광고 카드 없음", cards};
}
""")

        if not isinstance(data, dict) or not data.get("ok"):
            return outcome("판독불가", reason=(data or {}).get("reason", "PC DOM 판독 실패"))

        cards = []
        for rank, card in enumerate(data.get("cards", []), 1):
            item = dict(card)
            item["rank"] = rank
            item.pop("top", None)
            item.pop("left", None)
            cards.append(item)

        return filter_shopping_cards(cards)
    except Exception as error:
        return outcome("판독불가", reason=f"PC DOM 판독 오류: {type(error).__name__}")


def parse_powerlink(root, mobile):
    if mobile:
        links = [n for n in root.walk() if n.tag == "a" and re.search(r"a=pwl\.tit(?:&|$)", n.attrs.get("onclick") or "")]
        scope = links[0].closest("section") if links else None
    else:
        scope = next((n for n in root.walk() if n.cls("_pl_section")), None)
    if scope is None:
        return outcome("판독불가", reason="파워링크 영역 확인 실패")
    links = [n for n in scope.walk() if n.tag == "a" and ("a=pwl.tit" in (n.attrs.get("onclick") or "") or "plsearch" in (n.attrs.get("href") or ""))]
    nodes = list(dict.fromkeys(n.closest("li") for n in links if n.closest("li")))
    cards, seen = [], set()
    for node in nodes:
        aid = ad_id(node)
        if not aid:
            return outcome("판독불가", reason="파워링크 소재 ID 확인 실패")
        if aid in seen:
            continue
        seen.add(aid)
        children = list(node.walk())
        ranks = set()
        for n in children:
            onclick = n.attrs.get("onclick") or ""
            if "a=pwl.tit" in onclick:
                ranks.update(int(x) for x in re.findall(r"[?&]r=(\d+)", onclick))
        if mobile and len(ranks) != 1:
            return outcome("판독불가", reason="파워링크 순위 속성 오류")
        rank = next(iter(ranks)) if mobile else len(cards) + 1
        seller = next((clean(n.text()) for n in children if n.cls("site")), "")
        ours = seller == TARGET_SELLER or any(n.tag == "a" and (clean(n.text()) == TARGET_SELLER or store_match(n.attrs.get("href")) or store_match(n.attrs.get("onclick")) or store_match(n.text())) for n in children)
        cards.append({"rank": rank, "ad_id": aid, "product_id": "", "seller": TARGET_SELLER if ours else seller, "title": clean(node.text())[:100]})
    if not cards or [c["rank"] for c in cards] != list(range(1, len(cards) + 1)):
        return outcome("판독불가", reason="파워링크 순위 연속성 확인 실패")
    return outcome("확인", cards)


def display(result):
    if result["status"] != "확인":
        return result["status"]
    ranks = [c["rank"] for c in result["cards"] if c["seller"] == TARGET_SELLER]
    if ranks:
        return ", ".join(f"{r}위" for r in ranks)
    return "미노출(조회범위)" if all(c["seller"] for c in result["cards"]) else "판매자 확인불가"


class StopRun(RuntimeError):
    pass


def read_page_html(page, timeout_ms=15000):
    """이동 중 읽기 오류만 제한 시간 내 재시도. 새로고침/재검색하지 않음."""
    deadline = time.monotonic() + timeout_ms / 1000
    while time.monotonic() < deadline:
        if page.is_closed():
            raise StopRun("Chrome 탭이 닫혔습니다. 프로그램을 다시 실행하세요.")
        remaining = max(1, int((deadline - time.monotonic()) * 1000))
        before = page.url
        try:
            page.wait_for_load_state("domcontentloaded", timeout=min(2000, remaining))
            source = page.content()
            if page.url == before:
                return source
        except Exception as error:
            if page.is_closed():
                raise StopRun("Chrome 탭이 닫혔습니다.") from None
            message = str(error).lower()
            moving = any(text in message for text in (
                "page is navigating", "execution context was destroyed",
                "cannot find context with specified id", "cannot find execution context",
            ))
            if type(error).__name__ != "TimeoutError" and not moving:
                raise
        remaining = int((deadline - time.monotonic()) * 1000)
        if remaining <= 0:
            break
        page.wait_for_timeout(min(250, remaining))
    raise StopRun("페이지 이동이 계속되어 내용을 읽지 못했습니다. Chrome 상태를 확인한 뒤 다시 실행하세요.")


def logged_in(context):
    try:
        cookies = context.cookies()
    except Exception:
        return False
    names = {
        cookie["name"]
        for cookie in cookies
        if (
            "naver.com" in (cookie.get("domain") or "").lower()
            and (
                cookie.get("expires", -1) in (-1, None)
                or cookie.get("expires", 0) > time.time()
            )
        )
    }
    return {"NID_AUT", "NID_SES"}.issubset(names)


def explicit_login_page(page, root):
    host = (urlsplit(page.url).hostname or "").lower()
    if host in ("nid.naver.com", "accounts.naver.com", "login.naver.com"):
        return True

    body = clean(root.text()).lower()
    phrases = (
        "로그인해 주세요",
        "로그인 해 주세요",
        "로그인하세요",
        "세션이 만료",
        "로그아웃되었습니다",
    )
    return any(phrase in body for phrase in phrases)


def ensure_session(page, root):
    if logged_in(page.context):
        return

    # 쿠키 저장소가 탐색 직후 잠깐 갱신되는 경우를 위한 짧은 재확인입니다.
    for delay in (200, 400, 800):
        page.wait_for_timeout(delay)
        if logged_in(page.context):
            return

    if explicit_login_page(page, root):
        raise StopRun(
            "네이버 로그인 화면으로 이동했습니다. "
            "로그인 후 다시 실행하세요."
        )

    raise StopRun(
        "네이버 로그인 쿠키를 확인하지 못했습니다. "
        "보호조치/세션 만료 여부를 확인한 뒤 다시 실행하세요."
    )


def security_check(root):
    body = clean(root.text())
    words = ("보호조치", "보호 조치", "보안 확인", "보안퀴즈", "보안 퀴즈", "자동입력 방지", "자동 입력 방지", "비정상적인 접근", "비정상적인 요청", "접근이 제한", "접근 제한", "captcha")
    if any(word in body.lower() for word in words):
        raise StopRun("네이버 보호/보안 확인 감지. 해제 후 수동으로 다시 실행하세요.")


def wait_for_login(page, context, timeout_ms=20000):
    """로그인 쿠키와 네이버 홈 복귀를 함께 확인. 보호 화면은 즉시 중단."""
    print("로그인 후 페이지 이동이 끝나기를 기다리는 중...")
    deadline = time.monotonic() + timeout_ms / 1000
    previous_url, ready_since = None, None
    while time.monotonic() < deadline:
        remaining = max(1, int((deadline - time.monotonic()) * 1000))
        root = Document(read_page_html(page, timeout_ms=remaining)).root
        security_check(root)
        url = page.url
        home = urlsplit(url).hostname in ("naver.com", "www.naver.com", "m.naver.com")
        if home and clean(root.text()) and logged_in(context):
            if url != previous_url or ready_since is None:
                ready_since = time.monotonic()
            elif time.monotonic() - ready_since >= 0.75:
                return
        else:
            ready_since = None
        previous_url = url
        remaining = int((deadline - time.monotonic()) * 1000)
        if remaining <= 0:
            break
        page.wait_for_timeout(min(250, remaining))
    raise StopRun("로그인 완료/네이버 홈 복귀를 확인하지 못했습니다. 추가 인증 여부를 확인한 뒤 다시 실행하세요.")


def save_debug(page, keyword, device):
    safe = re.sub(r'[\\/:*?"<>|]', "_", keyword)
    path = DEBUG_DIR / f"{safe}_{device}_{datetime.now():%Y%m%d_%H%M%S_%f}"
    DEBUG_DIR.mkdir(parents=True, exist_ok=True)
    try:
        path.with_suffix(".html").write_text(read_page_html(page, timeout_ms=3000), encoding="utf-8")
        page.screenshot(path=str(path.with_suffix(".png")), full_page=False, timeout=5000)
        print("  진단 파일:", path.name)
    except Exception as error:
        print("  진단 저장 실패:", type(error).__name__)


def scan(page, keyword, mobile):
    host, where = ("m.search.naver.com", "m") if mobile else ("search.naver.com", "nexearch")
    url = f"https://{host}/search.naver?where={where}&query={quote(keyword)}"
    try:
        response = page.goto(url, wait_until="domcontentloaded", timeout=30000)
        if response and response.status in (401, 403, 429):
            raise StopRun(f"검색 접근 제한 HTTP {response.status}")
        if response and response.status >= 400:
            raise RuntimeError("검색 서버 오류")
        if urlsplit(page.url).hostname != host:
            raise StopRun("검색 페이지에서 다른 주소로 이동했습니다.")
        page.wait_for_timeout(PAGE_SETTLE_MS)
        for attempt in range(2):
            root = Document(read_page_html(page)).root
            security_check(root)
            if urlsplit(page.url).hostname != host:
                raise StopRun("검색 내용을 읽는 도중 다른 주소로 이동했습니다.")
            ensure_session(page, root)
            shopping_result = parse_shopping(root, mobile)
            if not mobile and shopping_result["status"] != "확인":
                live_result = parse_shopping_pc_live(page)
                if live_result["status"] == "확인":
                    shopping_result = live_result
            results = [shopping_result, parse_powerlink(root, mobile)]
            if all(r["status"] == "확인" for r in results) or attempt == 1:
                break
            page.wait_for_timeout(600)  # 재검색 없이 현재 DOM만 한 번 더 확인
        if any(r["status"] != "확인" for r in results):
            save_debug(page, keyword, "Mobile" if mobile else "PC")
        for result in results:
            result["observed_at"] = time.time()
        return results
    except StopRun:
        save_debug(page, keyword, "Mobile" if mobile else "PC")
        raise
    except Exception as error:
        save_debug(page, keyword, "Mobile" if mobile else "PC")
        return [dict(outcome("조회실패", reason=type(error).__name__), observed_at=time.time()) for _ in range(2)]


def _prepare_mobile_page(context, pc):
    """같은 로그인 쿠키를 공유하는 Mobile 탭을 1회만 생성합니다.

    키워드마다 PC/Mobile 상태를 오가지 않고, Mobile 배치 시작 시 한 번만
    모바일 렌더링 환경을 설정합니다. 스텔스/지문 위장은 사용하지 않습니다.
    """
    mobile = context.new_page()
    cdp = context.new_cdp_session(mobile)
    version = re.search(r"Chrome/[\d.]+", pc.evaluate("navigator.userAgent"))
    if not version:
        raise StopRun("Chrome 버전 확인 실패")
    cdp.send("Network.setUserAgentOverride", {
        "userAgent": "Mozilla/5.0 (Linux; Android 14; SM-S918N) AppleWebKit/537.36 (KHTML, like Gecko) " + version.group() + " Mobile Safari/537.36",
        "platform": "Android", "acceptLanguage": "ko-KR,ko;q=0.9",
    })
    cdp.send("Emulation.setDeviceMetricsOverride", {
        "width": 412, "height": 915, "deviceScaleFactor": 1, "mobile": True
    })
    cdp.send("Emulation.setTouchEmulationEnabled", {"enabled": True, "maxTouchPoints": 5})
    return mobile


def _batch_pause(index, total, device):
    if index >= total:
        return
    time.sleep(BETWEEN_KEYWORD_SECONDS)


def _connect_existing_chrome(p):
    """사용자가 먼저 띄운 전용 Chrome(remote debugging)에 연결합니다.

    일반적으로 이미 실행 중인 평소 Chrome에는 사후 연결할 수 없으므로
    동봉된 '01_start_login_chrome.bat'으로 연 Chrome을 사용합니다.
    """
    try:
        browser = p.chromium.connect_over_cdp(CDP_ENDPOINT)
    except Exception as error:
        raise StopRun(
            "열려 있는 전용 Chrome에 연결하지 못했습니다. "
            "먼저 01_start_login_chrome.bat을 실행하고 네이버 로그인 후 다시 시도하세요."
        ) from error
    contexts = browser.contexts
    if not contexts:
        raise StopRun("연결된 Chrome의 브라우저 컨텍스트를 찾지 못했습니다.")
    context = contexts[0]
    return browser, context


def _ensure_login_ready(pc, context):
    # 이미 네이버 페이지가 열려 있어도 홈에서 세션 상태를 한 번 명확히 확인합니다.
    response = pc.goto("https://www.naver.com/", wait_until="domcontentloaded", timeout=30000)
    if response and response.status >= 400:
        raise StopRun(f"네이버 홈 접근 오류 HTTP {response.status}")
    security_check(Document(read_page_html(pc)).root)
    if not logged_in(context):
        print("열린 Chrome에서 직접 로그인하세요. 비밀번호는 프로그램에 입력하지 마세요.")
        input("로그인 후 네이버 홈이 보이면 이 CMD에서 Enter > ")
    wait_for_login(pc, context)


def _launch_saved_state_browser(p):
    if not AUTH_STATE_FILE.exists():
        raise StopRun("저장된 네이버 로그인 상태가 없습니다. 로그인 상태 파일을 먼저 등록하세요.")
    launch_args = {"headless": True}
    try:
        browser = p.chromium.launch(**launch_args)
    except Exception as error:
        # Streamlit Cloud/GitHub 배포 직후 Chromium 바이너리가 아직 없는 경우 1회 설치합니다.
        if "Executable doesn't exist" not in str(error) and "executable" not in str(error).lower():
            raise
        print("Playwright Chromium이 없어 설치합니다. 최초 1회만 수행됩니다.")
        subprocess.run([sys.executable, "-m", "playwright", "install", "chromium"], check=True)
        browser = p.chromium.launch(**launch_args)
    context = browser.new_context(
        storage_state=str(AUTH_STATE_FILE),
        viewport={"width": 1280, "height": 900},
        locale="ko-KR",
    )
    return browser, context


def collect(keywords):
    from playwright.sync_api import sync_playwright

    # 키워드별 dict를 먼저 만들어 두면 PC/Mobile 배치를 분리해도 최종 출력 순서는 유지됩니다.
    by_keyword = {k: {"keyword": k, "surfaces": {}} for k in keywords}
    with sync_playwright() as p:
        remote_browser = None
        managed_context = None
        saved_browser = None
        if BROWSER_MODE == "connect":
            print("브라우저 방식: 사용자가 먼저 열어 둔 전용 Chrome에 연결")
            remote_browser, context = _connect_existing_chrome(p)
        elif BROWSER_MODE == "state":
            print("브라우저 방식: 서버에 저장된 네이버 로그인 세션 재사용")
            saved_browser, context = _launch_saved_state_browser(p)
        else:
            print("브라우저 방식: 저장된 전용 Chrome 프로필을 프로그램이 실행")
            managed_context = p.chromium.launch_persistent_context(
                str(PROFILE_DIR), channel="chrome", headless=False,
                viewport={"width": 1280, "height": 900}, locale="ko-KR",
                chromium_sandbox=True,
            )
            context = managed_context

        try:
            context.set_default_timeout(5000)
            context.set_default_navigation_timeout(30000)
            pages = [page for page in context.pages if not page.is_closed()]
            pc = pages[0] if pages else context.new_page()
            _ensure_login_ready(pc, context)
            print("네이버 로그인 쿠키 확인 완료. PC/Mobile은 동일 로그인 세션을 공유합니다.")

            print(f"\n=== 1단계: PC 전체 조회 ({len(keywords)}개) ===")
            for index, keyword in enumerate(keywords, 1):
                print(f"\n[PC {index}/{len(keywords)}] {keyword}")
                shopping_pc, power_pc = scan(pc, keyword, False)
                print("  파워링크 PC:", display(power_pc), "/ 쇼핑검색 PC:", display(shopping_pc))
                by_keyword[keyword]["surfaces"]["파워링크 PC"] = power_pc
                by_keyword[keyword]["surfaces"]["쇼핑 PC"] = shopping_pc
                _batch_pause(index, len(keywords), "PC")

            print(f"\nPC 전체 조회 완료. Mobile 전환 전 {BETWEEN_DEVICE_BATCH_SECONDS:.0f}초 휴식합니다.")
            time.sleep(BETWEEN_DEVICE_BATCH_SECONDS)

            print(f"\n=== 2단계: Mobile 전체 조회 ({len(keywords)}개) ===")
            mobile = _prepare_mobile_page(context, pc)
            for index, keyword in enumerate(keywords, 1):
                print(f"\n[Mobile {index}/{len(keywords)}] {keyword}")
                shopping_mo, power_mo = scan(mobile, keyword, True)
                print("  파워링크 Mobile:", display(power_mo), "/ 쇼핑검색 Mobile:", display(shopping_mo))
                by_keyword[keyword]["surfaces"]["파워링크 Mobile"] = power_mo
                by_keyword[keyword]["surfaces"]["쇼핑 Mobile"] = shopping_mo
                _batch_pause(index, len(keywords), "Mobile")

            return [by_keyword[k] for k in keywords]

        except (StopRun, KeyboardInterrupt):
            partial = [row for row in by_keyword.values() if row["surfaces"]]
            if partial:
                save_json(RESULT_DIR / f"중단결과_{datetime.now():%Y%m%d_%H%M%S}.json", partial)
            raise
        finally:
            # connect 모드는 사용자가 열어 둔 Chrome을 닫지 않습니다.
            # managed 모드만 프로그램이 띄운 Chrome을 종료합니다.
            if managed_context is not None:
                managed_context.close()
            if saved_browser is not None:
                saved_browser.close()
            # remote_browser.close()는 실제 Chrome을 종료할 수 있으므로 호출하지 않습니다.


# 기존 소재ID 엑셀: 쇼핑=grp ID, 파워링크=nkw ID. 소재 입찰 사용 여부는 바꾸지 않음.
def load_mapping():
    from openpyxl import load_workbook
    wb = load_workbook(MAPPING_FILE, read_only=True, data_only=True)
    try:
        for sheet in wb:
            rows = iter(sheet.values)
            headers = [clean(v) for v in next(rows, ())]
            if not all(h in headers for h in ("키워드",) + SURFACES):
                continue
            mapping = {}
            for row in rows:
                keyword = clean(row[headers.index("키워드")])
                if not keyword:
                    continue
                if keyword in mapping:
                    raise ValueError("매핑 엑셀에 중복 키워드: " + keyword)
                mapping[keyword] = {h: clean(row[headers.index(h)]) for h in SURFACES}
            return mapping
        raise ValueError("매핑 엑셀 컬럼을 확인하세요: 키워드 / " + " / ".join(SURFACES))
    finally:
        wb.close()


class NoRedirect(HTTPRedirectHandler):
    def redirect_request(self, *args, **kwargs):
        return None


class API:
    def __init__(self):
        if CONFIG_FILE.exists():
            self.config = json.loads(CONFIG_FILE.read_text(encoding="utf-8-sig"))
        else:
            self.config = {
                "customer_id": os.environ.get("NAVER_SEARCHAD_CUSTOMER_ID", ""),
                "api_key": os.environ.get("NAVER_SEARCHAD_API_KEY", ""),
                "secret_key": os.environ.get("NAVER_SEARCHAD_SECRET_KEY", ""),
            }
        for key in ("customer_id", "api_key", "secret_key"):
            if not clean(self.config.get(key)):
                raise ValueError("API 설정 누락: " + key)
        self.opener = build_opener(NoRedirect())
        self.last = 0

    def call(self, method, path, params=None, payload=None):
        time.sleep(max(0, 0.2 - (time.monotonic() - self.last)))
        timestamp = str(int(time.time() * 1000))
        message = f"{timestamp}.{method}.{path}".encode()
        signature = base64.b64encode(hmac.new(self.config["secret_key"].encode(), message, hashlib.sha256).digest()).decode()
        headers = {"Content-Type": "application/json; charset=UTF-8", "X-Timestamp": timestamp,
                   "X-API-KEY": self.config["api_key"], "X-Customer": str(self.config["customer_id"]), "X-Signature": signature}
        url = "https://api.searchad.naver.com" + path + ("?" + urlencode(params) if params else "")
        data = json.dumps(payload).encode() if payload is not None else None
        self.last = time.monotonic()
        try:
            with self.opener.open(Request(url, data=data, headers=headers, method=method), timeout=20) as response:
                body = response.read().decode()
                return json.loads(body) if body.strip() else None
        except HTTPError as error:
            raise StopRun(f"검색광고 API HTTP {error.code}. 설정/권한/요청을 확인하세요. 재시도하지 않습니다.") from None
        except (URLError, TimeoutError, OSError):
            raise StopRun("API 통신 실패. 변경 요청이었다면 적용 여부를 광고주센터에서 확인하세요.") from None


def active(obj):
    if not isinstance(obj, dict) or obj.get("userLock") or obj.get("delFlag") or obj.get("status") != "ELIGIBLE":
        raise ValueError("운영 가능 상태 아님 또는 상태 확인불가")


def adgroup_type_values(group):
    """Return the API's ad-group type fields in a normalized form.

    Search Ads responses have used both ``type`` and ``adgroupType`` in
    different endpoints/versions.  Keep both values so a conflicting response
    is still rejected instead of being silently treated as a match.
    """
    values = []
    for key in ("type", "adgroupType"):
        value = clean(group.get(key)).upper() if isinstance(group, dict) else ""
        if value and value not in values:
            values.append(value)
    return values


def target_state(api, surface, object_id, keyword):
    shopping = surface.startswith("쇼핑")
    prefix = "grp-" if shopping else "nkw-"
    if not re.fullmatch(prefix + r"[A-Za-z0-9-]+", object_id):
        raise ValueError("엑셀 ID 형식 오류: " + prefix + " 필요")
    path = ("/ncc/adgroups/" if shopping else "/ncc/keywords/") + object_id
    obj = api.call("GET", path)
    id_key = "nccAdgroupId" if shopping else "nccKeywordId"
    if not isinstance(obj, dict) or obj.get(id_key) != object_id:
        raise ValueError("API 대상 ID 불일치")
    active(obj)
    group = obj if shopping else api.call("GET", "/ncc/adgroups/" + obj["nccAdgroupId"])
    active(group)
    if group.get("nccAdgroupId") != obj.get("nccAdgroupId"):
        raise ValueError("소속 광고그룹 ID 불일치")
    observed_types = adgroup_type_values(group)
    # Naver's documented shopping enum is spelled SHOPING (one P).  Accept
    # the common SHOPPING spelling as a compatibility alias for older exports.
    expected_types = {"SHOPING", "SHOPPING"} if shopping else {"WEB_SITE"}
    if not observed_types or any(value not in expected_types for value in observed_types):
        shown = "/".join(observed_types) if observed_types else "확인불가"
        raise ValueError("광고그룹 유형 불일치: " + shown)
    if not shopping and clean(obj.get("keyword")) != keyword:
        raise ValueError("API 키워드와 조회 키워드 불일치")
    ads = api.call("GET", "/ncc/ads", {"nccAdgroupId": group["nccAdgroupId"]})
    if not isinstance(ads, list) or not ads:
        raise ValueError("그룹 소재 확인 실패")
    if any(a.get("nccAdgroupId") != group["nccAdgroupId"] for a in ads):
        raise ValueError("소재 소속 확인 실패")
    eligible = [a for a in ads if a.get("status") == "ELIGIBLE" and not a.get("userLock") and not a.get("delFlag")]
    if not eligible:
        raise ValueError("운영 가능 소재 없음")
    if shopping and any(a.get("type") != "SHOPPING_PRODUCT_AD" or a.get("adAttr", {}).get("useGroupBidAmt") is not True for a in eligible):
        raise ValueError("개별입찰/다른 유형 소재 포함: 그룹입찰 변경 제외")
    inherit = obj.get("useGroupBidAmt") if not shopping else None
    if not shopping and not isinstance(inherit, bool):
        raise ValueError("키워드 그룹입찰 사용 여부 확인 실패")
    current = int(group["bidAmt"] if inherit else obj["bidAmt"])
    minimum = 50 if shopping else 70
    if not minimum <= current <= MAX_BID:
        raise ValueError("현재 입찰가가 프로그램 허용 범위 밖: 수동 확인 필요")
    snapshot = {"object": obj, "group": group, "ads": sorted(ads, key=lambda a: a.get("nccAdId", ""))}
    fingerprint = hashlib.sha256(json.dumps(snapshot, sort_keys=True, ensure_ascii=False).encode()).hexdigest()
    return {"object": obj, "current": current, "fingerprint": fingerprint,
            "ad_ids": {a["nccAdId"] for a in eligible}, "minimum": minimum,
            "name": clean(group.get("name")), "inherit": inherit}


def matched_rank(result, ids):
    if result["status"] != "확인":
        raise ValueError(result["status"] + ": " + result.get("reason", ""))
    cards = result["cards"]
    matched = [c["rank"] for c in cards if c["ad_id"] in ids]
    if matched:
        return min(matched)
    if any(c.get("ad_id") in ids or (not c.get("ad_id") and c.get("seller") == TARGET_SELLER)
           for c in result.get("excluded_ad_plus", [])):
        raise ValueError("대상 광고가 광고+ 제외 항목에 포함: 미노출 인상 제외")
    if not cards:
        raise ValueError("소재를 특정할 수 없는 광고가 있어 미노출 확정 불가")
    # 화면에 잡힌 모든 광고의 판매자를 확인할 수 있고 대상 판매자가
    # 한 건도 없으면, 일부 경쟁 광고의 소재 ID가 비어 있어도 대상 소재는
    # 해당 1페이지(조회 범위)에 노출되지 않은 것으로 판단합니다.
    if all(c.get("seller") and c.get("seller") != TARGET_SELLER for c in cards):
        return 0
    if any(not c.get("ad_id") for c in cards):
        raise ValueError("대상 판매자 광고는 보이나 소재 ID를 특정할 수 없음")
    return 0  # 확인한 광고 영역 내 해당 그룹 소재 미노출


def new_bid(current, rank, minimum):
    if rank == 2:
        return current
    factor = Decimal("0.95") if rank == 1 else Decimal("1") if rank == 2 else Decimal("1.05")
    value = int((Decimal(current) * factor / 10).quantize(Decimal("1"), rounding=ROUND_HALF_UP)) * 10
    return min(MAX_BID, max(minimum, value))


def recent_ids():
    recent = set()
    if JOURNAL_FILE.exists():
        for line in JOURNAL_FILE.read_text(encoding="utf-8").splitlines():
            entry = json.loads(line)   # 이력이 손상됐으면 안전하게 중단
            if time.time() - entry["time"] < MIN_REBID_SECONDS:
                recent.add(entry["id"])
    return recent


def journal(item, status):
    entry = {"time": time.time(), "id": item["id"], "keyword": item["keyword"], "surface": item["surface"],
             "before": item["current"], "after": item["new"], "status": status}
    with JOURNAL_FILE.open("a", encoding="utf-8") as file:
        file.write(json.dumps(entry, ensure_ascii=False) + "\n")
        file.flush()
        os.fsync(file.fileno())


def build_plan(api, results, mapping, excluded, allow_absent=False):
    counts = Counter(v for row in mapping.values() for v in row.values() if v)
    recent, plan = recent_ids(), []
    for row in results:
        keyword = row["keyword"]
        if keyword in excluded:
            continue
        for surface in SURFACES:
            result = row["surfaces"][surface]
            oid = mapping.get(keyword, {}).get(surface, "")
            try:
                if not oid or counts[oid] != 1:
                    raise ValueError("ID 없음 또는 다른 항목과 ID 공유")
                if oid in recent:
                    raise ValueError("10분 내 변경 요청 이력 있음")
                if time.time() - result["observed_at"] > MAX_RESULT_AGE_SECONDS:
                    raise ValueError("조회 결과 30분 경과")
                if result["status"] != "확인":
                    raise ValueError(result["status"])
                state = target_state(api, surface, oid, keyword)
                rank = matched_rank(result, state["ad_ids"])
                if rank == 0 and not allow_absent:
                    raise ValueError("미노출 인상 미선택")
                proposed = new_bid(state["current"], rank, state["minimum"])
                item = {"keyword": keyword, "surface": surface, "id": oid, "name": state["name"],
                        "rank": rank, "current": state["current"], "new": proposed,
                        "fingerprint": state["fingerprint"], "observed_at": result["observed_at"]}
                plan.append(item)
                suffix = " (개별 키워드 입찰로 전환)" if state["inherit"] and proposed != state["current"] else ""
                print(f"{keyword} | {surface} | {rank or '미노출'} | {state['current']:,} → {proposed:,}원{suffix}")
            except ValueError as error:
                print(f"제외 | {keyword} | {surface} | {error}")
    return plan


def execute_plan(api, plan):
    global BID_REQUEST_ATTEMPTED
    seen = set()
    for item in plan:
        if item["new"] == item["current"]:
            continue
        oid = item["id"]
        if oid in seen or oid in recent_ids():
            raise StopRun("중복/최근 변경 감지. 추가 입찰을 중단했습니다.")
        seen.add(oid)
        if time.time() - item["observed_at"] > MAX_RESULT_AGE_SECONDS:
            raise StopRun("조회 결과가 오래되어 추가 입찰을 중단했습니다.")
        state = target_state(api, item["surface"], oid, item["keyword"])
        if state["fingerprint"] != item["fingerprint"]:
            raise StopRun("미리보기 후 광고 설정이 바뀌었습니다. 추가 입찰을 중단합니다.")
        if item["new"] != new_bid(state["current"], item["rank"], state["minimum"]):
            raise StopRun("변경 예정 금액 검증 실패")
        updated = dict(state["object"], bidAmt=item["new"])
        shopping = item["surface"].startswith("쇼핑")
        if not shopping:
            updated["useGroupBidAmt"] = False
        journal(item, "요청 전 기록")  # 통신 실패해도 다음 실행에서 바로 재입찰하지 않음
        path = "/ncc/adgroups/" + oid if shopping else "/ncc/keywords"
        try:
            BID_REQUEST_ATTEMPTED = True
            api.call("PUT", path, {"fields": "bidAmt"}, updated if shopping else [updated])
            check = api.call("GET", ("/ncc/adgroups/" if shopping else "/ncc/keywords/") + oid)
            key = "nccAdgroupId" if shopping else "nccKeywordId"
            if not isinstance(check, dict) or check.get(key) != oid or int(check.get("bidAmt", -1)) != item["new"] or (not shopping and check.get("useGroupBidAmt") is not False):
                raise StopRun("변경 후 값 확인 실패")
        except Exception as error:
            journal(item, "적용 여부 확인 필요")
            raise StopRun(f"{oid}: 추가 변경 중단. 적용 여부를 광고주센터에서 확인하세요. {error}") from None
        journal(item, "변경 확인 완료")
        print(f"변경 확인 | {item['keyword']} | {item['surface']} | {item['current']:,} → {item['new']:,}원")


def save_results(results, stamp):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    save_json(RESULT_DIR / f"순위상세_{stamp}.json", results)
    print("\n최종 결과 (각 항목은 해당 검색 시점의 순위)")
    print("키워드 | " + " | ".join(SURFACES))
    for row in results:
        print(row["keyword"] + " | " + " | ".join(display(row["surfaces"][s]) for s in SURFACES))
    wb = Workbook()
    ws = wb.active
    ws.title = "통합 순위"
    ws.append(["키워드"] + list(SURFACES))
    for row in results:
        ws.append([row["keyword"]] + [display(row["surfaces"][s]) for s in SURFACES])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for col in "ABCDE":
        ws.column_dimensions[col].width = 24
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    path = RESULT_DIR / f"황금이네_통합순위_{stamp}.xlsx"
    wb.save(path)
    print("\n결과 저장:", path)


def main():
    RESULT_DIR.mkdir(parents=True, exist_ok=True)
    print("황금이네 순위체커 + 승인형 5% 입찰관리 [PC 전체 → Mobile 전체 안정화본]")
    print("실행 파일:", BASE_DIR / "golden_integrated_checker.py")
    print("기준: 통합검색 첫 쇼핑 페이지의 광고 / 첫 파워링크 영역")
    print("PC/Mobile 화면 표시는 판매자 기준, 입찰은 엑셀 대상의 실제 소재 ID 기준입니다.")
    print(f"조회 간격: 키워드 {BETWEEN_KEYWORD_SECONDS:.0f}초 / PC→Mobile {BETWEEN_DEVICE_BATCH_SECONDS:.0f}초")
    print("보호/보안 확인 페이지가 감지되면 재시도·우회하지 않고 즉시 중단합니다.")
    raw = input(f"검사 키워드: Enter=전체 {len(KEYWORDS)}개 / 일부는 쉼표로 입력 > ").strip()
    keywords = list(dict.fromkeys(clean(k) for k in raw.split(",") if clean(k))) if raw else KEYWORDS
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    results = collect(keywords)
    if not results:
        return
    save_results(results, stamp)
    if not yes("현재 입찰가를 조회하고 변경 예정표를 만들까요?"):
        return
    mapping = load_mapping()
    excluded = {clean(k) for k in input("입찰 제외 키워드(쉼표 구분, 없으면 Enter) > ").split(",") if clean(k)}
    allow_absent = yes("해당 소재의 미노출을 확인한 항목도 +5% 인상할까요?")
    api = API()
    print("\n규칙: 1위 -5% / 2위 유지 / 3위 이하 +5% / 10원 반올림")
    print("쇼핑 최소 50원 / 파워링크 최소 70원 / 프로그램 상한 20,000원")
    print("쇼핑은 그룹 기본입찰가를 바꾸므로 그룹입찰을 사용하는 다른 소재에도 영향을 줍니다.")
    plan = build_plan(api, results, mapping, excluded, allow_absent)
    save_json(RESULT_DIR / f"입찰변경계획_{stamp}.json", plan)
    changes = [p for p in plan if p["new"] != p["current"]]
    print(f"\n실제 변경 대상 {len(changes)}개. 아직 입찰가를 변경하지 않았습니다.")
    if changes and input("대상과 금액 확인 후 '변경'을 정확히 입력 > ").strip() == "변경":
        execute_plan(api, changes)
    else:
        print("입찰가 변경 없이 종료합니다.")


def show_stop_notice():
    if BID_REQUEST_ATTEMPTED:
        print("입찰 변경을 시도한 항목이 있습니다. 적용 여부를 광고주센터와 이력에서 확인하세요.")
    else:
        print("이번 실행에서는 입찰 변경 요청을 보내지 않았습니다.")


if __name__ == "__main__":
    lock = None
    try:
        lock = acquire_lock()
        main()
    except KeyboardInterrupt:
        print("\n사용자 중단.")
        show_stop_notice()
    except Exception as error:
        print(f"\n중단: {error}")
        show_stop_notice()
    finally:
        if lock:
            lock.close()
        try:
            input("\nEnter를 누르면 종료합니다 > ")
        except (EOFError, KeyboardInterrupt):
            pass
